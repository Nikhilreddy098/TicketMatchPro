import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(file_path, report_title, test_cases):
    """
    Generates a production-grade multi-sheet Excel report using openpyxl.
    """
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    wb = openpyxl.Workbook()

    # Styles
    header_fill = PatternFill(start_color="6C3BFF", end_color="6C3BFF", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    pass_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid") # Light green
    pass_font = Font(name="Calibri", size=10, bold=True, color="065F46")

    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light red
    fail_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    skip_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light amber
    skip_font = Font(name="Calibri", size=10, bold=True, color="92400E")

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    columns = [
        "Test ID", "Module", "Test Name", "Status", "Execution Time (s)",
        "Priority", "Preconditions", "Test Steps", "Expected Result", "Actual Result"
    ]

    # Sheet 1: Executed Test Cases
    ws_all = wb.active
    ws_all.title = "Executed Test Cases"
    ws_all.append(columns)

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws_all.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    passed_list = []
    failed_list = []
    skipped_list = []

    for tc in test_cases:
        status = tc.get("status", "PASS").upper()
        row = [
            tc.get("test_id", ""),
            tc.get("module", ""),
            tc.get("test_name", ""),
            status,
            tc.get("duration", 0.05),
            tc.get("priority", "P1"),
            tc.get("preconditions", "N/A"),
            tc.get("test_steps", "Execute verification step"),
            tc.get("expected", "Success HTTP 200 / UI Match"),
            tc.get("actual", "Verified clean output")
        ]
        ws_all.append(row)
        current_row = ws_all.max_row

        # Format status cell
        status_cell = ws_all.cell(row=current_row, column=4)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        if status == "PASS":
            status_cell.fill = pass_fill
            status_cell.font = pass_font
            passed_list.append(tc)
        elif status == "FAIL":
            status_cell.fill = fail_fill
            status_cell.font = fail_font
            failed_list.append(tc)
        else:
            status_cell.fill = skip_fill
            status_cell.font = skip_font
            skipped_list.append(tc)

        for c_idx in range(1, len(columns) + 1):
            ws_all.cell(row=current_row, column=c_idx).border = thin_border

    # Sheet 2: Passed Tests
    ws_pass = wb.create_sheet(title="Passed Tests")
    ws_pass.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws_pass.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
    for tc in passed_list:
        ws_pass.append([
            tc.get("test_id"), tc.get("module"), tc.get("test_name"), "PASS",
            tc.get("duration", 0.05), tc.get("priority"), tc.get("preconditions"),
            tc.get("test_steps"), tc.get("expected"), tc.get("actual")
        ])

    # Sheet 3: Failed Tests
    ws_fail = wb.create_sheet(title="Failed Tests")
    ws_fail.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws_fail.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
    for tc in failed_list:
        ws_fail.append([
            tc.get("test_id"), tc.get("module"), tc.get("test_name"), "FAIL",
            tc.get("duration", 0.05), tc.get("priority"), tc.get("preconditions"),
            tc.get("test_steps"), tc.get("expected"), tc.get("actual")
        ])

    # Sheet 4: Skipped Tests
    ws_skip = wb.create_sheet(title="Skipped Tests")
    ws_skip.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws_skip.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
    for tc in skipped_list:
        ws_skip.append([
            tc.get("test_id"), tc.get("module"), tc.get("test_name"), "SKIPPED",
            tc.get("duration", 0.05), tc.get("priority"), tc.get("preconditions"),
            tc.get("test_steps"), tc.get("expected"), tc.get("actual")
        ])

    # Sheet 5: Execution Metrics
    ws_metrics = wb.create_sheet(title="Execution Metrics")
    total_count = len(test_cases)
    pass_count = len(passed_list)
    fail_count = len(failed_list)
    skip_count = len(skipped_list)
    pass_rate = round((pass_count / total_count * 100), 2) if total_count > 0 else 0.0

    ws_metrics.append(["Metric Name", "Metric Value"])
    ws_metrics.append(["Report Title", report_title])
    ws_metrics.append(["Total Executed Test Cases", total_count])
    ws_metrics.append(["Passed Test Cases", pass_count])
    ws_metrics.append(["Failed Test Cases", fail_count])
    ws_metrics.append(["Skipped Test Cases", skip_count])
    ws_metrics.append(["Pass Percentage (%)", f"{pass_rate}%"])

    for r in range(1, 8):
        ws_metrics.cell(row=r, column=1).font = Font(bold=True)

    # Sheet 6: Defect Summary
    ws_defects = wb.create_sheet(title="Defect Summary")
    ws_defects.append(["Defect ID", "Test ID", "Module", "Failure Description", "Severity"])
    for idx, tc in enumerate(failed_list, 1):
        ws_defects.append([
            f"DEF-{idx:03d}", tc.get("test_id"), tc.get("module"),
            tc.get("actual", "Unexpected behavior"), tc.get("priority", "High")
        ])

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(file_path)
    print(f"✅ Excel Report successfully generated: {file_path} ({len(test_cases)} Test Cases)")
